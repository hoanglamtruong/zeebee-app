"""Tạo 3 file Excel "Tải về dùng" của app.zeebee.vn.

Chạy:  python build_excel.py  → dist/*.xlsx
Nguyên tắc: font Arial; ô NHẬP tô vàng nhạt, ô CÔNG THỨC để trắng; mỗi file có sheet "Hướng dẫn"
và dòng ví dụ; chỉ dùng hàm thời Excel 2007 (SUMIFS, COUNTIFS, INDEX/MATCH, IFERROR, EOMONTH)
để mở được bằng Excel 2010+ và Google Sheets.
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / "dist"
OUT.mkdir(exist_ok=True)

FONT = "Arial"
NAVY = "0A2D47"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor="E8EEF4")
THIN = Side(style="thin", color="C9D3DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
VND = '#,##0;-#,##0;"-"'
DATE = "dd/mm/yyyy"


def font(**kw):
    return Font(name=FONT, size=kw.pop("size", 11), **kw)


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 32


def style_range(ws, rows, cols, *, fill=None, fmt=None, bold=False):
    for r in rows:
        for c in cols:
            cell = ws.cell(row=r, column=c)
            cell.font = font(bold=bold)
            cell.border = BORDER
            if fill:
                cell.fill = fill
            if fmt:
                cell.number_format = fmt


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(italic=True, color="5B6B80", size=10)


def guide_sheet(wb, name, lines):
    ws = wb.active
    ws.title = "Hướng dẫn"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110
    ws["A1"] = name
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Sản phẩm của Zee Bee · app.zeebee.vn — dùng được bằng Excel (2010 trở lên) hoặc Google Sheets."
    ws["A2"].font = font(italic=True, color="5B6B80")
    r = 4
    ws.cell(row=r, column=1, value="Ô tô VÀNG = ô bạn nhập. Ô TRẮNG = công thức tự tính, đừng xóa.").fill = INPUT_FILL
    ws.cell(row=r, column=1).font = font(bold=True)
    r += 2
    for line in lines:
        c = ws.cell(row=r, column=1, value=line)
        is_head = line and line[0].isdigit() and line[1] == "."
        c.font = font(bold=is_head, color=NAVY if is_head else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.cell(row=r + 1, column=1, value="Hỗ trợ: nhắn Zalo cho Seller đã bán sản phẩm cho bạn, hoặc Zalo 0976 094 453.").font = font(color="5B6B80")
    return ws


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def list_validation(ws, formula, cell_range, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorTitle="Giá trị không có trong danh sách", error="Chọn một mục trong danh sách.")
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Gợi ý", prompt, True
    ws.add_data_validation(dv)
    dv.add(cell_range)


# ════════════════════════════════════════════════════════════════════════════
# 1. SỔ THU CHI GIA ĐÌNH
# ════════════════════════════════════════════════════════════════════════════
def so_thu_chi():
    wb = Workbook()
    guide_sheet(wb, "Sổ thu chi gia đình", [
        "1. Bắt đầu",
        "   Sheet \"Danh mục\": sửa tên hạng mục Thu / Chi cho hợp với gia đình bạn (thêm vào ô vàng còn trống).",
        "   Sheet \"Tổng hợp\": nhập NĂM muốn xem ở ô vàng.",
        "2. Ghi thu chi hằng ngày",
        "   Sheet \"Nhật ký\": mỗi khoản một dòng — Ngày, Loại (Thu/Chi), Hạng mục (chọn từ danh sách), Số tiền, Ghi chú.",
        "   3 dòng ví dụ ở đầu: xóa đi khi bắt đầu ghi thật. Sổ chứa được 2.000 dòng.",
        "3. Xem báo cáo",
        "   Sheet \"Tổng hợp\": Thu, Chi, Còn lại theo từng tháng + bảng chi theo hạng mục × tháng — tự cập nhật.",
        "4. Lưu ý",
        "   Không sửa ô trắng (công thức). Muốn thêm dòng nhật ký: ghi tiếp xuống, đừng chèn cột.",
    ])

    # Danh mục
    dm = wb.create_sheet("Danh mục")
    title(dm, "Danh mục hạng mục", "Sửa / thêm tên hạng mục trong ô vàng.")
    header(dm, 4, ["Hạng mục THU", "Hạng mục CHI"], [28, 28])
    thu = ["Lương", "Thưởng", "Kinh doanh", "Lãi tiết kiệm", "Được cho / tặng", "Thu khác"]
    chi = ["Ăn uống", "Nhà ở / thuê nhà", "Điện, nước, internet", "Đi lại, xăng xe", "Con cái, học hành",
           "Sức khỏe", "Mua sắm", "Giải trí, du lịch", "Hiếu hỉ, biếu tặng", "Tiết kiệm, đầu tư", "Trả nợ", "Chi khác"]
    for i in range(20):
        dm.cell(row=5 + i, column=1, value=thu[i] if i < len(thu) else None)
        dm.cell(row=5 + i, column=2, value=chi[i] if i < len(chi) else None)
    style_range(dm, range(5, 25), range(1, 3), fill=INPUT_FILL)
    add_name(wb, "DS_THU", "'Danh mục'!$A$5:$A$24")
    add_name(wb, "DS_CHI", "'Danh mục'!$B$5:$B$24")
    # Danh sách chọn (data validation) chỉ nhận vùng 1 cột → cột D gộp Thu rồi Chi, tự cập nhật.
    dm.cell(row=4, column=4, value="Danh sách gộp (tự động)")
    dm["D4"].font, dm["D4"].fill, dm["D4"].border = font(bold=True, color="FFFFFF"), HEAD_FILL, BORDER
    dm.column_dimensions["D"].width = 28
    for i in range(20):
        dm.cell(row=5 + i, column=4, value=f'=IF(A{5 + i}="","",A{5 + i})')
        dm.cell(row=25 + i, column=4, value=f'=IF(B{5 + i}="","",B{5 + i})')
    style_range(dm, range(5, 45), [4])
    add_name(wb, "DS_HANGMUC", "'Danh mục'!$D$5:$D$44")

    # Nhật ký
    nk = wb.create_sheet("Nhật ký")
    title(nk, "Nhật ký thu chi", "Mỗi khoản một dòng. Cột Tháng/Năm tự tính.")
    header(nk, 4, ["Ngày", "Loại", "Hạng mục", "Số tiền (đ)", "Ghi chú", "Tháng", "Năm"], [13, 9, 24, 15, 34, 8, 8])
    first, last = 5, 2004
    today = dt.date.today()
    examples = [
        (dt.date(today.year, today.month, 1), "Thu", "Lương", 15000000, "Lương tháng (ví dụ — xóa dòng này)"),
        (dt.date(today.year, today.month, 2), "Chi", "Ăn uống", 250000, "Đi chợ (ví dụ — xóa dòng này)"),
        (dt.date(today.year, today.month, 3), "Chi", "Điện, nước, internet", 820000, "Tiền điện (ví dụ — xóa dòng này)"),
    ]
    for r in range(first, last + 1):
        nk.cell(row=r, column=6, value=f'=IF(A{r}="","",MONTH(A{r}))')
        nk.cell(row=r, column=7, value=f'=IF(A{r}="","",YEAR(A{r}))')
    for i, (d, loai, hm, tien, note) in enumerate(examples):
        r = first + i
        nk.cell(row=r, column=1, value=d)
        nk.cell(row=r, column=2, value=loai)
        nk.cell(row=r, column=3, value=hm)
        nk.cell(row=r, column=4, value=tien)
        nk.cell(row=r, column=5, value=note)
    style_range(nk, range(first, last + 1), range(1, 6), fill=INPUT_FILL)
    style_range(nk, range(first, last + 1), range(6, 8))
    for r in range(first, last + 1):
        nk.cell(row=r, column=1).number_format = DATE
        nk.cell(row=r, column=4).number_format = VND
    list_validation(nk, '"Thu,Chi"', f"B{first}:B{last}")
    list_validation(nk, "=DS_HANGMUC", f"C{first}:C{last}", "Chọn hạng mục; sửa danh sách ở sheet Danh mục.")
    nk.freeze_panes = "A5"

    # Tổng hợp
    th = wb.create_sheet("Tổng hợp")
    title(th, "Tổng hợp theo tháng", "Nhập năm ở ô vàng. Mọi số liệu tự tính từ Nhật ký.")
    th["A3"], th["B3"] = "Năm:", today.year
    th["A3"].font = font(bold=True)
    th["B3"].fill, th["B3"].font, th["B3"].border = INPUT_FILL, font(bold=True), BORDER
    header(th, 5, ["Tháng", "Tổng thu", "Tổng chi", "Còn lại", "Tích lũy từ đầu năm"], [12, 16, 16, 16, 20])
    NK = "'Nhật ký'"
    rng = lambda col: f"{NK}!${col}${first}:${col}${last}"
    for m in range(1, 13):
        r = 5 + m
        th.cell(row=r, column=1, value=m)
        th.cell(row=r, column=2, value=f'=SUMIFS({rng("D")},{rng("B")},"Thu",{rng("F")},A{r},{rng("G")},$B$3)')
        th.cell(row=r, column=3, value=f'=SUMIFS({rng("D")},{rng("B")},"Chi",{rng("F")},A{r},{rng("G")},$B$3)')
        th.cell(row=r, column=4, value=f"=B{r}-C{r}")
        th.cell(row=r, column=5, value=f"=SUM($D$6:D{r})")
    style_range(th, range(6, 18), range(1, 6), fmt=VND)
    for r in range(6, 18):
        th.cell(row=r, column=1).number_format = '"Tháng "0'
    th["A18"] = "Cả năm"
    for c, col in ((2, "B"), (3, "C"), (4, "D")):
        th.cell(row=18, column=c, value=f"=SUM({col}6:{col}17)")
    style_range(th, [18], range(1, 6), fill=TOTAL_FILL, fmt=VND, bold=True)

    # Chi theo hạng mục × tháng
    top = 21
    th.cell(row=top - 1, column=1, value="Chi theo hạng mục").font = font(bold=True, color=NAVY, size=12)
    header(th, top, ["Hạng mục chi"] + [f"T{m}" for m in range(1, 13)] + ["Cả năm"])
    th.column_dimensions["A"].width = 24
    for c in range(2, 15):
        th.column_dimensions[get_column_letter(c)].width = max(th.column_dimensions[get_column_letter(c)].width or 0, 12)
    for i in range(20):
        r = top + 1 + i
        th.cell(row=r, column=1, value=f"=IF('Danh mục'!B{5 + i}=\"\",\"\",'Danh mục'!B{5 + i})")
        for m in range(1, 13):
            th.cell(row=r, column=1 + m,
                    value=f'=IF($A{r}="","",SUMIFS({rng("D")},{rng("B")},"Chi",{rng("C")},$A{r},{rng("F")},{m},{rng("G")},$B$3))')
        th.cell(row=r, column=14, value=f'=IF($A{r}="","",SUM(B{r}:M{r}))')
    style_range(th, range(top + 1, top + 21), range(1, 15), fmt=VND)
    th.freeze_panes = "B6"

    path = OUT / "So-thu-chi-gia-dinh.xlsx"
    wb.save(path)
    return path


# ════════════════════════════════════════════════════════════════════════════
# 2. QUẢN LÝ BÁN HÀNG SHOP NHỎ
# ════════════════════════════════════════════════════════════════════════════
def ban_hang():
    wb = Workbook()
    guide_sheet(wb, "Quản lý bán hàng shop nhỏ", [
        "1. Khai báo sản phẩm",
        "   Sheet \"Sản phẩm\": mỗi sản phẩm một dòng — Mã SP (ngắn, không trùng), Tên, Đơn vị, Giá vốn, Giá bán, Tồn đầu kỳ,",
        "   Mức cảnh báo (tồn bằng hoặc dưới mức này sẽ hiện \"Sắp hết\"). Cột Nhập / Bán / Tồn hiện tại tự tính.",
        "2. Nhập hàng",
        "   Sheet \"Nhập kho\": Ngày, Mã SP (chọn), Số lượng, Đơn giá nhập. Tên và thành tiền tự hiện.",
        "3. Bán hàng",
        "   Sheet \"Bán hàng\": Ngày, Mã đơn, Khách hàng, Mã SP, Số lượng, Giảm giá (nếu có).",
        "   Giá bán và giá vốn tự lấy từ sheet Sản phẩm; Doanh thu, Tiền vốn, Lãi gộp tự tính.",
        "   Đơn nhiều sản phẩm: ghi mỗi sản phẩm một dòng, cùng Mã đơn.",
        "4. Báo cáo",
        "   Sheet \"Báo cáo\": nhập Năm ở ô vàng — xem Doanh thu, Tiền vốn, Lãi gộp, Số lượng bán theo tháng.",
        "5. Lưu ý",
        "   Dòng ví dụ (sản phẩm SP01, SP02 và các dòng nhập/bán mẫu): xóa đi khi bắt đầu dùng thật.",
        "   Đổi giá bán ở sheet Sản phẩm sẽ ảnh hưởng các dòng bán cũ — muốn giữ giá cũ, ghi vào cột \"Giá bán khác\".",
    ])
    today = dt.date.today()

    sp = wb.create_sheet("Sản phẩm")
    title(sp, "Danh mục sản phẩm & tồn kho", "Ô vàng: khai báo. Ô trắng: tự tính.")
    header(sp, 4, ["Mã SP", "Tên sản phẩm", "Đơn vị", "Giá vốn (đ)", "Giá bán (đ)", "Tồn đầu kỳ",
                   "Mức cảnh báo", "Đã nhập", "Đã bán", "Tồn hiện tại", "Trạng thái"],
           [10, 30, 9, 13, 13, 11, 11, 10, 10, 12, 13])
    sp_first, sp_last = 5, 304
    nk_first, nk_last = 5, 2004
    bh_first, bh_last = 5, 5004
    for r in range(sp_first, sp_last + 1):
        sp.cell(row=r, column=8, value=f"=IF(A{r}=\"\",\"\",SUMIFS('Nhập kho'!$D${nk_first}:$D${nk_last},'Nhập kho'!$B${nk_first}:$B${nk_last},A{r}))")
        sp.cell(row=r, column=9, value=f"=IF(A{r}=\"\",\"\",SUMIFS('Bán hàng'!$F${bh_first}:$F${bh_last},'Bán hàng'!$D${bh_first}:$D${bh_last},A{r}))")
        sp.cell(row=r, column=10, value=f'=IF(A{r}="","",F{r}+H{r}-I{r})')
        sp.cell(row=r, column=11, value=f'=IF(A{r}="","",IF(J{r}<=0,"Hết hàng",IF(J{r}<=G{r},"Sắp hết","Còn hàng")))')
    for i, row in enumerate([("SP01", "Bình gốm men mờ (ví dụ)", "cái", 150000, 280000, 10, 3),
                             ("SP02", "Kệ gỗ treo tường (ví dụ)", "cái", 220000, 450000, 5, 2)]):
        for c, v in enumerate(row, start=1):
            sp.cell(row=sp_first + i, column=c, value=v)
    style_range(sp, range(sp_first, sp_last + 1), range(1, 8), fill=INPUT_FILL)
    style_range(sp, range(sp_first, sp_last + 1), range(8, 12))
    for r in range(sp_first, sp_last + 1):
        for c in (4, 5):
            sp.cell(row=r, column=c).number_format = VND
        for c in (6, 7, 8, 9, 10):
            sp.cell(row=r, column=c).number_format = "#,##0"
    sp.freeze_panes = "C5"
    add_name(wb, "DS_MASP", f"'Sản phẩm'!$A${sp_first}:$A${sp_last}")
    SPR = lambda col: f"'Sản phẩm'!${col}${sp_first}:${col}${sp_last}"
    lookup = lambda key, col: f"IFERROR(INDEX({SPR(col)},MATCH({key},{SPR('A')},0)),\"\")"

    nk = wb.create_sheet("Nhập kho")
    title(nk, "Phiếu nhập kho", "Mỗi lần nhập một sản phẩm một dòng.")
    header(nk, 4, ["Ngày", "Mã SP", "Tên sản phẩm", "Số lượng", "Đơn giá nhập (đ)", "Thành tiền (đ)", "Nhà cung cấp / ghi chú",
                   "Tháng", "Năm"],
           [13, 10, 30, 10, 15, 16, 30, 8, 8])
    for r in range(nk_first, nk_last + 1):
        nk.cell(row=r, column=3, value=f'=IF(B{r}="","",{lookup(f"B{r}", "B")})')
        nk.cell(row=r, column=6, value=f'=IF(OR(D{r}="",E{r}=""),"",D{r}*E{r})')
        nk.cell(row=r, column=8, value=f'=IF(A{r}="","",MONTH(A{r}))')
        nk.cell(row=r, column=9, value=f'=IF(A{r}="","",YEAR(A{r}))')
    nk.cell(row=nk_first, column=1, value=dt.date(today.year, today.month, 1))
    nk.cell(row=nk_first, column=2, value="SP01")
    nk.cell(row=nk_first, column=4, value=6)
    nk.cell(row=nk_first, column=5, value=150000)
    nk.cell(row=nk_first, column=7, value="Ví dụ — xóa dòng này")
    style_range(nk, range(nk_first, nk_last + 1), [1, 2, 4, 5, 7], fill=INPUT_FILL)
    style_range(nk, range(nk_first, nk_last + 1), [3, 6, 8, 9])
    for r in range(nk_first, nk_last + 1):
        nk.cell(row=r, column=1).number_format = DATE
        nk.cell(row=r, column=4).number_format = "#,##0"
        nk.cell(row=r, column=5).number_format = VND
        nk.cell(row=r, column=6).number_format = VND
    list_validation(nk, "=DS_MASP", f"B{nk_first}:B{nk_last}")
    nk.freeze_panes = "A5"

    bh = wb.create_sheet("Bán hàng")
    title(bh, "Sổ bán hàng", "Đơn nhiều sản phẩm: mỗi sản phẩm một dòng, cùng Mã đơn.")
    header(bh, 4, ["Ngày", "Mã đơn", "Khách hàng", "Mã SP", "Tên sản phẩm", "Số lượng", "Giá bán khác (nếu có)",
                   "Giảm giá (đ)", "Giá bán áp dụng", "Doanh thu (đ)", "Tiền vốn (đ)", "Lãi gộp (đ)", "Tháng", "Năm"],
           [13, 11, 20, 10, 28, 10, 14, 12, 14, 15, 15, 15, 8, 8])
    for r in range(bh_first, bh_last + 1):
        bh.cell(row=r, column=5, value=f'=IF(D{r}="","",{lookup(f"D{r}", "B")})')
        bh.cell(row=r, column=9, value=f'=IF(D{r}="","",IF(G{r}<>"",G{r},{lookup(f"D{r}", "E")}))')
        bh.cell(row=r, column=10, value=f'=IF(OR(D{r}="",F{r}=""),"",F{r}*I{r}-N(H{r}))')
        bh.cell(row=r, column=11, value=f'=IF(OR(D{r}="",F{r}=""),"",F{r}*{lookup(f"D{r}", "D")})')
        bh.cell(row=r, column=12, value=f'=IF(J{r}="","",J{r}-K{r})')
        bh.cell(row=r, column=13, value=f'=IF(A{r}="","",MONTH(A{r}))')
        bh.cell(row=r, column=14, value=f'=IF(A{r}="","",YEAR(A{r}))')
    for i, row in enumerate([(dt.date(today.year, today.month, 2), "DH001", "Chị Lan (ví dụ)", "SP01", 2, None, 20000),
                             (dt.date(today.year, today.month, 2), "DH001", "Chị Lan (ví dụ)", "SP02", 1, None, None),
                             (dt.date(today.year, today.month, 5), "DH002", "Anh Minh (ví dụ)", "SP01", 1, 260000, None)]):
        r = bh_first + i
        for c, v in zip((1, 2, 3, 4, 6, 7, 8), row):
            bh.cell(row=r, column=c, value=v)
    style_range(bh, range(bh_first, bh_last + 1), [1, 2, 3, 4, 6, 7, 8], fill=INPUT_FILL)
    style_range(bh, range(bh_first, bh_last + 1), [5, 9, 10, 11, 12, 13, 14])
    for r in range(bh_first, bh_last + 1):
        bh.cell(row=r, column=1).number_format = DATE
        bh.cell(row=r, column=6).number_format = "#,##0"
        for c in (7, 8, 9, 10, 11, 12):
            bh.cell(row=r, column=c).number_format = VND
    list_validation(bh, "=DS_MASP", f"D{bh_first}:D{bh_last}")
    bh.freeze_panes = "A5"

    bc = wb.create_sheet("Báo cáo")
    title(bc, "Báo cáo kinh doanh theo tháng", "Nhập năm ở ô vàng.")
    bc["A3"], bc["B3"] = "Năm:", today.year
    bc["A3"].font = font(bold=True)
    bc["B3"].fill, bc["B3"].font, bc["B3"].border = INPUT_FILL, font(bold=True), BORDER
    header(bc, 5, ["Tháng", "Doanh thu (đ)", "Tiền vốn (đ)", "Lãi gộp (đ)", "Tỷ lệ lãi gộp", "Số lượng bán", "Tiền nhập hàng (đ)"],
           [12, 16, 16, 16, 13, 13, 18])
    B = lambda col: f"'Bán hàng'!${col}${bh_first}:${col}${bh_last}"
    for m in range(1, 13):
        r = 5 + m
        bc.cell(row=r, column=1, value=m)
        bc.cell(row=r, column=2, value=f"=SUMIFS({B('J')},{B('M')},A{r},{B('N')},$B$3)")
        bc.cell(row=r, column=3, value=f"=SUMIFS({B('K')},{B('M')},A{r},{B('N')},$B$3)")
        bc.cell(row=r, column=4, value=f"=B{r}-C{r}")
        bc.cell(row=r, column=5, value=f'=IF(B{r}=0,"",D{r}/B{r})')
        bc.cell(row=r, column=6, value=f"=SUMIFS({B('F')},{B('M')},A{r},{B('N')},$B$3)")
        NKR = lambda col: f"'Nhập kho'!${col}${nk_first}:${col}${nk_last}"
        bc.cell(row=r, column=7, value=f"=SUMIFS({NKR('F')},{NKR('H')},A{r},{NKR('I')},$B$3)")
    style_range(bc, range(6, 18), range(1, 8), fmt=VND)
    for r in range(6, 18):
        bc.cell(row=r, column=1).number_format = '"Tháng "0'
        bc.cell(row=r, column=5).number_format = "0.0%"
        bc.cell(row=r, column=6).number_format = "#,##0"
    bc["A18"] = "Cả năm"
    for c in (2, 3, 4, 6, 7):
        col = get_column_letter(c)
        bc.cell(row=18, column=c, value=f"=SUM({col}6:{col}17)")
    bc["E18"] = '=IF(B18=0,"",D18/B18)'
    style_range(bc, [18], range(1, 8), fill=TOTAL_FILL, fmt=VND, bold=True)
    bc["E18"].number_format = "0.0%"
    bc["F18"].number_format = "#,##0"

    path = OUT / "Quan-ly-ban-hang-shop-nho.xlsx"
    wb.save(path)
    return path


# ════════════════════════════════════════════════════════════════════════════
# 3. CHẤM CÔNG & TÍNH LƯƠNG
# ════════════════════════════════════════════════════════════════════════════
def cham_cong():
    wb = Workbook()
    guide_sheet(wb, "Chấm công & tính lương", [
        "1. Cài đặt (làm mỗi tháng)",
        "   Sheet \"Cài đặt\": nhập Tháng, Năm cần chấm công; Số ngày công chuẩn (thường 26), Số giờ/ngày, Hệ số tăng ca.",
        "2. Khai báo nhân viên",
        "   Sheet \"Nhân viên\": Mã NV, Họ tên, Chức vụ, Lương tháng (đủ công), Phụ cấp cố định.",
        "3. Chấm công hằng ngày",
        "   Sheet \"Chấm công\": chọn Mã NV ở cột A; mỗi ngày gõ 1 ký hiệu:",
        "      X = đi làm đủ ngày (1 công)    N = nửa ngày (0,5 công)    P = nghỉ phép CÓ lương (1 công)",
        "      O = nghỉ KHÔNG lương (0 công)  để trống = chưa chấm / không làm",
        "   Cột \"Giờ tăng ca\": tổng số giờ tăng ca trong tháng.",
        "4. Bảng lương",
        "   Sheet \"Bảng lương\": tự tính lương theo công, tiền tăng ca, phụ cấp. Nhập thêm Thưởng và Khấu trừ (tạm ứng,",
        "   BHXH, phạt…) ở ô vàng → ra Thực nhận.",
        "5. Lưu ý",
        "   File KHÔNG tự tính BHXH và thuế TNCN — nếu có, nhập số tiền vào cột Khấu trừ.",
        "   Sang tháng mới: lưu một bản sao file (Save As) rồi xóa ký hiệu chấm công cũ, đổi Tháng ở sheet Cài đặt.",
        "   Dòng ví dụ (NV01, NV02): xóa hoặc sửa thành nhân viên thật.",
    ])
    today = dt.date.today()

    cd = wb.create_sheet("Cài đặt")
    title(cd, "Cài đặt kỳ chấm công")
    rows = [("Tháng", today.month, "1–12"), ("Năm", today.year, ""), ("Số ngày công chuẩn / tháng", 26, "Đủ công = nhận đủ lương tháng"),
            ("Số giờ làm / ngày", 8, "Dùng tính lương 1 giờ"), ("Hệ số tăng ca", 1.5, "1,5 = 150% lương giờ")]
    cd.column_dimensions["A"].width = 30
    cd.column_dimensions["B"].width = 12
    cd.column_dimensions["C"].width = 36
    for i, (k, v, note) in enumerate(rows):
        r = 4 + i
        cd.cell(row=r, column=1, value=k).font = font(bold=True)
        c = cd.cell(row=r, column=2, value=v)
        c.fill, c.font, c.border = INPUT_FILL, font(), BORDER
        cd.cell(row=r, column=3, value=note).font = font(italic=True, color="5B6B80", size=10)
    for nm, ref in (("CD_THANG", "$B$4"), ("CD_NAM", "$B$5"), ("CD_CONG_CHUAN", "$B$6"), ("CD_GIO_NGAY", "$B$7"), ("CD_HE_SO_TC", "$B$8")):
        add_name(wb, nm, f"'Cài đặt'!{ref}")

    nv = wb.create_sheet("Nhân viên")
    title(nv, "Danh sách nhân viên")
    header(nv, 4, ["Mã NV", "Họ tên", "Chức vụ", "Lương tháng (đ)", "Phụ cấp (đ)"], [10, 26, 18, 16, 14])
    nv_first, nv_last = 5, 104
    for i, row in enumerate([("NV01", "Nguyễn Văn A (ví dụ)", "Phục vụ", 6000000, 500000),
                             ("NV02", "Trần Thị B (ví dụ)", "Thu ngân", 7000000, 300000)]):
        for c, v in enumerate(row, start=1):
            nv.cell(row=nv_first + i, column=c, value=v)
    style_range(nv, range(nv_first, nv_last + 1), range(1, 6), fill=INPUT_FILL)
    for r in range(nv_first, nv_last + 1):
        nv.cell(row=r, column=4).number_format = VND
        nv.cell(row=r, column=5).number_format = VND
    add_name(wb, "DS_MANV", f"'Nhân viên'!$A${nv_first}:$A${nv_last}")
    NVR = lambda col: f"'Nhân viên'!${col}${nv_first}:${col}${nv_last}"
    nvlookup = lambda key, col: f"IFERROR(INDEX({NVR(col)},MATCH({key},{NVR('A')},0)),\"\")"

    cc = wb.create_sheet("Chấm công")
    cc["A1"] = '="Bảng chấm công tháng "&CD_THANG&"/"&CD_NAM'
    cc["A1"].font = font(bold=True, size=14, color=NAVY)
    cc["A2"] = "X = đủ ngày · N = nửa ngày · P = phép có lương · O = nghỉ không lương"
    cc["A2"].font = font(italic=True, color="5B6B80", size=10)
    # Hàng 4: số ngày; hàng 5: thứ trong tuần. Ngày vượt quá số ngày của tháng để trống.
    cc.cell(row=4, column=1, value="Mã NV")
    cc.cell(row=4, column=2, value="Họ tên")
    cc.column_dimensions["A"].width = 9
    cc.column_dimensions["B"].width = 24
    day_col0 = 3  # cột C = ngày 1
    for d in range(1, 32):
        col = day_col0 + d - 1
        L = get_column_letter(col)
        cc.cell(row=4, column=col, value=f"=IF({d}>DAY(EOMONTH(DATE(CD_NAM,CD_THANG,1),0)),\"\",{d})")
        cc.cell(row=5, column=col, value=f'=IF({L}4="","",CHOOSE(WEEKDAY(DATE(CD_NAM,CD_THANG,{d})),"CN","T2","T3","T4","T5","T6","T7"))')
        cc.column_dimensions[L].width = 4.2
    last_day_col = day_col0 + 30
    tot_cols = ["Công", "Phép", "Nghỉ KL", "Giờ tăng ca"]
    for i, t in enumerate(tot_cols):
        cc.cell(row=4, column=last_day_col + 1 + i, value=t)
        cc.column_dimensions[get_column_letter(last_day_col + 1 + i)].width = 10
    for c in range(1, last_day_col + 5):
        for r in (4, 5):
            cell = cc.cell(row=r, column=c)
            cell.font = font(bold=True, color="FFFFFF", size=10)
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER
    cc_first, cc_last = 6, 105
    dfirst, dlast = get_column_letter(day_col0), get_column_letter(last_day_col)
    cong, phep, nkl, gio = (get_column_letter(last_day_col + i) for i in range(1, 5))
    for r in range(cc_first, cc_last + 1):
        cc.cell(row=r, column=2, value=f'=IF(A{r}="","",{nvlookup(f"A{r}", "B")})')
        rr = f"{dfirst}{r}:{dlast}{r}"
        cc[f"{cong}{r}"] = f'=IF(A{r}="","",COUNTIF({rr},"X")+COUNTIF({rr},"P")+0.5*COUNTIF({rr},"N"))'
        cc[f"{phep}{r}"] = f'=IF(A{r}="","",COUNTIF({rr},"P"))'
        cc[f"{nkl}{r}"] = f'=IF(A{r}="","",COUNTIF({rr},"O"))'
    # ví dụ
    import itertools
    pattern = itertools.cycle(["X"] * 5 + ["N", ""])
    for i, ma in enumerate(["NV01", "NV02"]):
        r = cc_first + i
        cc.cell(row=r, column=1, value=ma)
        for d in range(1, 11):
            cc.cell(row=r, column=day_col0 + d - 1, value=next(pattern) if i == 0 else ("P" if d == 4 else "X"))
        cc[f"{gio}{r}"] = 4 if i == 0 else 0
    style_range(cc, range(cc_first, cc_last + 1), [1] + list(range(day_col0, last_day_col + 1)) + [last_day_col + 4], fill=INPUT_FILL)
    style_range(cc, range(cc_first, cc_last + 1), [2] + list(range(last_day_col + 1, last_day_col + 4)))
    for r in range(cc_first, cc_last + 1):
        for c in range(day_col0, last_day_col + 1):
            cc.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    list_validation(cc, "=DS_MANV", f"A{cc_first}:A{cc_last}")
    list_validation(cc, '"X,N,P,O"', f"{dfirst}{cc_first}:{dlast}{cc_last}")
    cc.freeze_panes = f"{dfirst}{cc_first}"

    bl = wb.create_sheet("Bảng lương")
    bl["A1"] = '="Bảng lương tháng "&CD_THANG&"/"&CD_NAM'
    bl["A1"].font = font(bold=True, size=14, color=NAVY)
    bl["A2"] = "Ô vàng: Thưởng, Khấu trừ. Còn lại tự tính từ Chấm công và Nhân viên."
    bl["A2"].font = font(italic=True, color="5B6B80", size=10)
    header(bl, 4, ["Mã NV", "Họ tên", "Lương tháng", "Công thực tế", "Lương theo công", "Giờ tăng ca", "Tiền tăng ca",
                   "Phụ cấp", "Thưởng", "Khấu trừ", "Thực nhận", "Ký nhận"],
           [9, 24, 14, 11, 15, 10, 14, 13, 13, 13, 15, 14])
    CC = lambda col: f"'Chấm công'!${col}${cc_first}:${col}${cc_last}"
    cclookup = lambda key, col: f"IFERROR(INDEX({CC(col)},MATCH({key},{CC('A')},0)),0)"
    for i in range(cc_last - cc_first + 1):
        r = 5 + i
        src = cc_first + i
        bl.cell(row=r, column=1, value=f"=IF('Chấm công'!A{src}=\"\",\"\",'Chấm công'!A{src})")
        bl.cell(row=r, column=2, value=f'=IF(A{r}="","",{nvlookup(f"A{r}", "B")})')
        bl.cell(row=r, column=3, value=f'=IF(A{r}="","",N({nvlookup(f"A{r}", "D")}))')
        bl.cell(row=r, column=4, value=f'=IF(A{r}="","",N({cclookup(f"A{r}", cong)}))')
        bl.cell(row=r, column=5, value=f'=IF(A{r}="","",ROUND(C{r}/CD_CONG_CHUAN*D{r},0))')
        bl.cell(row=r, column=6, value=f'=IF(A{r}="","",N({cclookup(f"A{r}", gio)}))')
        bl.cell(row=r, column=7, value=f'=IF(A{r}="","",ROUND(C{r}/CD_CONG_CHUAN/CD_GIO_NGAY*CD_HE_SO_TC*F{r},0))')
        bl.cell(row=r, column=8, value=f'=IF(A{r}="","",N({nvlookup(f"A{r}", "E")}))')
        bl.cell(row=r, column=11, value=f'=IF(A{r}="","",E{r}+G{r}+H{r}+N(I{r})-N(J{r}))')
    bl_last = 5 + (cc_last - cc_first)
    style_range(bl, range(5, bl_last + 1), range(1, 13), fmt=VND)
    style_range(bl, range(5, bl_last + 1), [9, 10], fill=INPUT_FILL, fmt=VND)
    for r in range(5, bl_last + 1):
        bl.cell(row=r, column=4).number_format = "0.0"
        bl.cell(row=r, column=6).number_format = "0.0"
        bl.cell(row=r, column=1).number_format = "General"
        bl.cell(row=r, column=2).number_format = "General"
    tr = bl_last + 1
    bl.cell(row=tr, column=2, value="TỔNG CỘNG")
    for c in (5, 7, 8, 9, 10, 11):
        col = get_column_letter(c)
        bl.cell(row=tr, column=c, value=f"=SUM({col}5:{col}{bl_last})")
    style_range(bl, [tr], range(1, 13), fill=TOTAL_FILL, fmt=VND, bold=True)
    bl.freeze_panes = "C5"

    path = OUT / "Cham-cong-tinh-luong.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    for fn in (so_thu_chi, ban_hang, cham_cong):
        print("OK", fn())
